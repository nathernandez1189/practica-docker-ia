"""Genera versiones Word y PDF del informe Markdown y sus evidencias."""
from pathlib import Path
import re
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, Preformatted
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.colors import HexColor as _HexColor, white
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/INFORME.md'
INK, TEAL = '122D43', '096C70'
def HexColor(value):
    return _HexColor('#' + value.lstrip('#'))
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.8)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.color.rgb = RGBColor.from_string(INK)
doc.core_properties.title = 'Práctica de contenedores Docker e IA'
doc.core_properties.author = 'Natalia Hernández'
footer = section.footer.paragraphs[0]
footer.text = 'Natalia Hernández · Práctica Docker e IA · '
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyLab', fontName='Helvetica', fontSize=9.5, leading=14, spaceAfter=7, textColor=HexColor(INK)))
styles.add(ParagraphStyle(name='TitleLab', fontName='Helvetica-Bold', fontSize=25, leading=30, spaceAfter=20, textColor=HexColor(INK)))
styles.add(ParagraphStyle(name='HeadingLab', fontName='Helvetica-Bold', fontSize=14, leading=18, spaceBefore=14, spaceAfter=8, textColor=HexColor(TEAL), keepWithNext=True))
styles.add(ParagraphStyle(name='CellLab', fontName='Helvetica', fontSize=8.2, leading=11, textColor=HexColor(INK)))
styles.add(ParagraphStyle(name='CaptionLab', fontName='Helvetica-Oblique', fontSize=8, leading=11, spaceAfter=10, textColor=HexColor(TEAL)))
story = []

def plain(text):
    return re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text).replace('**', '').replace('`', '').replace('*', '')

def rich(text):
    text = escape(text)
    text = re.sub(r'\[([^]]+)\]\((https?://[^)]+)\)', r'<link href="\2" color="#096c70">\1</link>', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'<b>\1</b>', text)
    text = re.sub(r'`([^`]+)`', r'<font name="Courier">\1</font>', text)
    text = re.sub(r'\*([^*]+)\*', r'<i>\1</i>', text)
    return text

def para(text, bullet=False):
    p = doc.add_paragraph(style='List Bullet' if bullet else None)
    for i, part in enumerate(re.split(r'(\*\*[^*]+\*\*)', plain(text) if '**' not in text else text)):
        run = p.add_run(plain(part))
        run.bold = part.startswith('**')
    story.append(Paragraph(('&#8226; ' if bullet else '') + rich(text), styles['BodyLab']))

lines = SOURCE.read_text().splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith('```'):
        code = []
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph()
        run = p.add_run('\n'.join(code)); run.font.name = 'Consolas'; run.font.size = Pt(8)
        story.append(Preformatted('\n'.join(code), ParagraphStyle('CodeLab', fontName='Courier', fontSize=7.5, leading=11, backColor=HexColor('F0F4F2'), borderPadding=8, spaceAfter=12), maxLineLength=96))
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', c) for c in cells): rows.append(cells)
            i += 1
        i -= 1
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = 'Light Shading Accent 1'
        for j, cell in enumerate(rows[0]): table.rows[0].cells[j].text = plain(cell)
        for row in rows[1:]:
            cells = table.add_row().cells
            for j, cell in enumerate(row): cells[j].text = plain(cell)
        pdfrows = [[Paragraph(rich(c), styles['CellLab']) for c in row] for row in rows]
        widths = [160, 340] if len(rows[0]) == 2 else [500 / len(rows[0])] * len(rows[0])
        table_pdf = Table(pdfrows, colWidths=widths, repeatRows=1, hAlign='LEFT')
        table_pdf.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('DEEBE5')),('ROWBACKGROUNDS',(0,1),(-1,-1),[white,HexColor('F5F7F4')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('LINEBELOW',(0,0),(-1,0),.7,HexColor(TEAL))]))
        story.extend([table_pdf, Spacer(1,10)])
        doc.add_paragraph()
    elif line.startswith('!['):
        match = re.match(r'!\[([^]]*)\]\(([^)]+)\)', line)
        caption, relative = match.groups()
        path = (SOURCE.parent / relative).resolve()
        if not path.exists(): raise FileNotFoundError(path)
        doc.add_picture(str(path), width=Inches(6.5))
        doc.add_paragraph(caption, style='Caption')
        w, h = PILImage.open(path).size
        scale = min(500/w, 360/h)
        story.extend([Image(str(path), width=w*scale, height=h*scale), Paragraph(escape(caption), styles['CaptionLab'])])
    elif line.startswith('# '):
        doc.add_heading(line[2:], 0)
        story.append(Paragraph(rich(line[2:]), styles['TitleLab']))
    elif line.startswith('## '):
        doc.add_heading(line[3:], 1)
        story.append(Paragraph(rich(line[3:]), styles['HeadingLab']))
    else:
        para(line[2:] if line.startswith('- ') else line, bullet=line.startswith('- '))
    i += 1

doc.save(ROOT / 'docs/Informe-Practica-Docker-IA.docx')

def footer_pdf(canvas, document):
    canvas.saveState()
    canvas.setStrokeColor(HexColor('CDD8CD'))
    canvas.line(47, 40, A4[0]-47, 40)
    canvas.setFont('Helvetica', 8)
    canvas.setFillColor(HexColor(TEAL))
    canvas.drawString(47, 27, 'Natalia Hernández · Práctica Docker e IA')
    canvas.drawRightString(A4[0]-47, 27, str(document.page))
    canvas.restoreState()

pdf = SimpleDocTemplate(str(ROOT / 'docs/Informe-Practica-Docker-IA.pdf'), pagesize=A4, rightMargin=47, leftMargin=47, topMargin=45, bottomMargin=57, title='Práctica Docker e IA', author='Natalia Hernández')
pdf.build(story, onFirstPage=footer_pdf, onLaterPages=footer_pdf)
print('Informe Word y PDF generados.')
